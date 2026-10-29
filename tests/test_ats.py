import glob
import io
import os
import pytest
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml

from app.ats import ats_check
from app import extract as X


def make_pdf(stream_content: str) -> bytes:
    stream_bytes = stream_content.encode("latin1")
    pdf = f"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >> >> >> >> endobj
4 0 obj << /Length {len(stream_bytes)} >> stream
{stream_content}
endstream endobj
xref
0 5
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000300 00000 n 
trailer << /Size 5 /Root 1 0 R >>
startxref
450
%%EOF""".encode("latin1")
    return pdf


def test_clean_resume_is_low():
    body = (
        "Jane Doe\n"
        "Software Engineer\n"
        "Email: jane.doe@example.com | Phone: 555-019-2834\n\n"
        "Education\n"
        "B.S. in Computer Science, Tech University, 2024\n\n"
        "Work Experience\n"
        "Software Engineer Intern at Acme Corp\n"
        "- Built microservices using Python and Docker.\n"
        "- Optimized SQL database queries reducing latency by 35%.\n\n"
        "Projects\n"
        "Distributed KV Store: Implemented Raft consensus in Python and Go.\n\n"
        "Skills\n"
        "Python, Docker, SQL, Git, Linux\n"
    )
    clean_stream = (
        "BT 0 0 0 rg /F1 11 Tf\n"
        "50 700 Td (Jane Doe - Software Engineer) Tj\n"
        "0 -15 Td (Email: jane.doe@example.com | Phone: 555-019-2834) Tj\n"
        "0 -20 Td (Education) Tj\n"
        "0 -15 Td (B.S. in Computer Science, Tech University, 2024) Tj\n"
        "0 -20 Td (Work Experience) Tj\n"
        "0 -15 Td (Software Engineer Intern at Acme Corp) Tj\n"
        "0 -15 Td (Built microservices using Python and Docker.) Tj\n"
        "0 -20 Td (Projects) Tj\n"
        "0 -15 Td (Distributed KV Store in Python and Go) Tj\n"
        "0 -20 Td (Skills) Tj\n"
        "0 -15 Td (Python, Docker, SQL, Git, Linux) Tj\n"
        "ET"
    )
    pdf_bytes = make_pdf(clean_stream)
    res = ats_check(pdf_bytes, "clean.pdf", body)
    assert res["level"] == "Low", f"Expected Low for clean resume, got {res['level']} (score: {res['risk_score']})"
    assert res["risk_score"] < 15


def test_white_text_synthetic_is_high():
    body = (
        "Candidate Cheater\n"
        "Email: cheater@example.com\n"
        "Education: Some University\n"
        "Experience: Engineer\n"
    )
    # Visible text + White hidden text keywords
    stream = (
        "BT 0 0 0 rg /F1 11 Tf 50 700 Td (Candidate Cheater Education Experience) Tj ET\n"
        "BT 1 1 1 rg /F1 10 Tf 50 500 Td (python java kubernetes aws docker react typescript microservices) Tj ET"
    )
    pdf_bytes = make_pdf(stream)
    # Include zero-width chars or repeated keywords to verify High risk
    text_with_cheats = body + " \u200b\u200c\u200d\ufeff hidden white keywords python aws docker react"
    res = ats_check(pdf_bytes, "white_text.pdf", text_with_cheats)
    assert res["level"] == "High"
    finding_ids = [f["id"] for f in res["findings"]]
    assert "hidden_white_text" in finding_ids
    white_finding = next(f for f in res["findings"] if f["id"] == "hidden_white_text")
    assert white_finding["severity"] == "high"
    assert "fix" in white_finding and len(white_finding["fix"]) > 10


def test_tiny_text_synthetic_is_high():
    body = (
        "Candidate Cheater\n"
        "Email: tiny@example.com\n"
        "Education: Tech College\n"
    )
    stream = (
        "BT 0 0 0 rg /F1 11 Tf 50 700 Td (Candidate Cheater Education) Tj ET\n"
        "BT 0 0 0 rg /F1 3 Tf 50 400 Td (microscopic tiny keywords python react kubernetes docker cloud aws) Tj ET"
    )
    pdf_bytes = make_pdf(stream)
    text_with_cheats = body + " \u200b\u200c\u200d\ufeff microscopic tiny stuffed keywords"
    res = ats_check(pdf_bytes, "tiny_text.pdf", text_with_cheats)
    assert res["level"] == "High"
    finding_ids = [f["id"] for f in res["findings"]]
    assert "hidden_tiny_text" in finding_ids
    tiny_finding = next(f for f in res["findings"] if f["id"] == "hidden_tiny_text")
    assert tiny_finding["severity"] == "high"


def test_zero_width_chars_synthetic_is_high():
    # Resume laced with invisible zero-width spaces across keywords and missing sections
    body = (
        "Candidate Stealth\n"
        "P\u200by\u200ct\u200dh\u200eo\u200fn K\u200bu\u200cb\u200de\u200fr\u200fn\u200be\u200ct\u200de\u200fs "
        "A\u200bW\u200bS D\u200bo\u200bc\u200bk\u200be\u200fr G\u200bo\u200bl\u200ba\u200bn\u200bg"
    )
    pdf_bytes = make_pdf("BT 0 0 0 rg /F1 11 Tf 50 700 Td (Candidate Stealth) Tj ET")
    res = ats_check(pdf_bytes, "zero_width.pdf", body)
    assert res["level"] == "High"
    finding_ids = [f["id"] for f in res["findings"]]
    assert "invisible_zero_width_chars" in finding_ids


def test_docx_white_and_vanish_synthetic():
    doc = Document()
    doc.add_paragraph("Alice Developer | alice@dev.local | 555-010-9988")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Software Engineering, 2023")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Developer at Startup Inc")

    # Add paragraph with white font
    p_white = doc.add_paragraph()
    r_white = p_white.add_run("hidden white text keywords python aws docker")
    r_white.font.color.rgb = RGBColor(255, 255, 255)

    # Add run with w:vanish
    p_vanish = doc.add_paragraph()
    r_vanish = p_vanish.add_run("vanished secret keywords")
    vanish_xml = parse_xml(r'<w:vanish xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    r_vanish._r.get_or_add_rPr().append(vanish_xml)

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = ats_check(docx_bytes, "cheater.docx", "Alice Developer Education Experience hidden white text keywords python aws docker vanished secret keywords")
    finding_ids = [f["id"] for f in res["findings"]]
    assert "hidden_white_text_docx" in finding_ids
    assert "hidden_vanish_text_docx" in finding_ids


def test_all_four_sample_resumes_are_not_high():
    resume_files = sorted(glob.glob("seed/resumes/*.pdf")) or sorted(glob.glob("data/resumes/*.pdf"))

    assert len(resume_files) == 4, f"Expected 4 sample resumes, found {len(resume_files)}"

    for fpath in resume_files:
        name = os.path.basename(fpath)
        with open(fpath, "rb") as f:
            raw = f.read()
        txt = X.read_text(raw, name)
        skills = X.extract(raw, "resume", name, use_cache=True).get("skills", [])

        res = ats_check(raw, name, txt, skills_extracted=skills)
        assert res["level"] != "High", (
            f"Sample resume {name} failed: evaluated to {res['level']} (score: {res['risk_score']})! "
            f"Findings: {[f['title'] for f in res['findings']]}"
        )
