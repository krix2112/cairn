import io
import re
import unicodedata
from typing import Dict, List, Optional, Tuple, Any

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

from pypdf import PdfReader
from docx import Document

SEVERITY_WEIGHTS = {
    "high": 25,
    "medium": 12,
    "low": 5,
    "info": 0
}

STANDARD_SECTIONS = {
    "education": r"\b(education|academics|qualifications|university|college|degree)\b",
    "experience": r"\b(experience|work experience|employment|internships?|professional experience)\b",
    "projects": r"\b(projects|personal projects|technical projects|academic projects)\b",
    "skills": r"\b(skills|technical skills|competencies|technologies|tools)\b"
}

EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b")
ZERO_WIDTH_RE = re.compile(r"[\u200B-\u200D\uFEFF\u200E\u200F\u202A-\u202E\u00AD\u2060]")
EXCESSIVE_PADDING_RE = re.compile(r"(?:[ \t]{25,}|\n{6,})")


def _is_color_white(color) -> bool:
    """Check if color is white or near-white (channels >= 0.93 or >= 237)."""
    if color is None:
        return False
    if isinstance(color, (int, float)):
        val = float(color)
        return val >= 0.93 if val <= 1.0 else val >= 237
    if isinstance(color, (list, tuple)):
        if len(color) == 3: # RGB
            return all(
                (float(c) >= 0.93 if float(c) <= 1.0 else float(c) >= 237)
                for c in color
            )
        elif len(color) == 4: # CMYK
            return all(
                (float(c) <= 0.07 if float(c) <= 1.0 else float(c) <= 18)
                for c in color
            )
        elif len(color) == 1: # Grayscale
            val = float(color[0])
            return val >= 0.93 if val <= 1.0 else val >= 237
    return False


def _check_hidden_pdf(pdf_bytes: bytes) -> List[Dict[str, Any]]:
    findings = []
    if not pdfplumber or not pdf_bytes:
        return findings
    white_chars = []
    tiny_chars = []
    oob_chars = []

    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            for p_idx, page in enumerate(pdf.pages):
                width, height = float(page.width), float(page.height)
                chars = page.chars or []
                for c in chars:
                    t = c.get("text", "")
                    if not t.strip():
                        continue
                    
                    # 1. White / near-white fill
                    n_col = c.get("non_stroking_color")
                    s_col = c.get("stroking_color")
                    if _is_color_white(n_col) or _is_color_white(s_col):
                        white_chars.append(t)

                    # 2. Font size < 5 pt
                    sz = float(c.get("size", 10.0))
                    if sz < 5.0 and sz > 0.0:
                        tiny_chars.append(t)

                    # 3. Outside page box
                    x0 = float(c.get("x0", 0.0))
                    y0 = float(c.get("top", 0.0))
                    x1 = float(c.get("x1", 0.0))
                    y1 = float(c.get("bottom", 0.0))
                    if x0 < -2.0 or y0 < -2.0 or x1 > width + 2.0 or y1 > height + 2.0:
                        oob_chars.append(t)
    except Exception:
        pass

    if len(white_chars) >= 3:
        sample = "".join(white_chars[:60])
        findings.append({
            "id": "hidden_white_text",
            "severity": "high",
            "title": "Hidden White/Camouflage Text Detected",
            "detail": f"Detected {len(white_chars)} characters rendered in white or near-white color fill (>= 0.93 threshold).",
            "evidence": f"Text sample: '{sample}'",
            "fix": "Remove white-colored keywords. Modern ATS parsers extract all raw text streams and flag camouflage text as deceptive manipulation."
        })

    if len(tiny_chars) >= 5:
        sample = "".join(tiny_chars[:60])
        findings.append({
            "id": "hidden_tiny_text",
            "severity": "high",
            "title": "Microscopic Text Detected (< 5 pt)",
            "detail": f"Detected {len(tiny_chars)} characters with font size below 5 pt.",
            "evidence": f"Text sample: '{sample}'",
            "fix": "Remove microscopic text. Text below 5 pt is commonly flagged as keyword stuffing by enterprise screening systems."
        })

    if len(oob_chars) >= 5:
        sample = "".join(oob_chars[:60])
        findings.append({
            "id": "hidden_offpage_text",
            "severity": "high",
            "title": "Off-Page Text Detected",
            "detail": f"Detected {len(oob_chars)} characters positioned outside the printable page boundary.",
            "evidence": f"Text sample: '{sample}'",
            "fix": "Keep all content within standard page margins (0.5 to 1 inch). Position coordinates outside the page box indicate hidden text."
        })

    return findings


def _check_hidden_docx(docx_bytes: bytes) -> List[Dict[str, Any]]:
    findings = []
    white_runs = []
    vanish_runs = []
    tiny_runs = []

    try:
        doc = Document(io.BytesIO(docx_bytes))
        # Scan main paragraphs
        for p in doc.paragraphs:
            for r in p.runs:
                txt = r.text
                if not txt or not txt.strip():
                    continue

                # 1. w:vanish
                vanish_elem = r._element.xpath('.//w:vanish')
                if vanish_elem:
                    vanish_runs.append(txt)

                # 2. Font color near white
                color_elem = r._element.xpath('.//w:color')
                if color_elem:
                    val = color_elem[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                    if val.upper() in ('FFFFFF', 'FEFEFE', 'FDFDFD', 'FAFAFA', 'F5F5F5'):
                        white_runs.append(txt)

                # 3. Font size < 5 pt (w:sz < 10)
                sz_elem = r._element.xpath('.//w:sz')
                if sz_elem:
                    try:
                        sz_val = int(sz_elem[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '20'))
                        if sz_val < 10:
                            tiny_runs.append(txt)
                    except ValueError:
                        pass

        # Scan tables for hidden formatting
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            txt = r.text
                            if not txt or not txt.strip():
                                continue
                            if r._element.xpath('.//w:vanish'):
                                vanish_runs.append(txt)
                            color_elem = r._element.xpath('.//w:color')
                            if color_elem:
                                val = color_elem[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                                if val.upper() in ('FFFFFF', 'FEFEFE', 'FDFDFD', 'FAFAFA', 'F5F5F5'):
                                    white_runs.append(txt)
                            sz_elem = r._element.xpath('.//w:sz')
                            if sz_elem:
                                try:
                                    if int(sz_elem[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '20')) < 10:
                                        tiny_runs.append(txt)
                                except ValueError:
                                    pass
    except Exception:
        pass

    if white_runs:
        findings.append({
            "id": "hidden_white_text_docx",
            "severity": "high",
            "title": "Hidden White Text in DOCX",
            "detail": f"Found {len(white_runs)} text runs formatted with white font fill.",
            "evidence": f"Sample: '{white_runs[0][:50]}'",
            "fix": "Remove runs with white font color in your document."
        })
    if vanish_runs:
        findings.append({
            "id": "hidden_vanish_text_docx",
            "severity": "high",
            "title": "Hidden Text Tag (w:vanish) in DOCX",
            "detail": f"Found {len(vanish_runs)} hidden text runs formatted with Word vanish property.",
            "evidence": f"Sample: '{vanish_runs[0][:50]}'",
            "fix": "Unhide or delete hidden text elements in Word before submitting."
        })
    if tiny_runs:
        findings.append({
            "id": "hidden_tiny_text_docx",
            "severity": "high",
            "title": "Microscopic Text in DOCX (< 5 pt)",
            "detail": f"Found {len(tiny_runs)} runs with font size smaller than 5 pt.",
            "evidence": f"Sample: '{tiny_runs[0][:50]}'",
            "fix": "Maintain standard readable font sizes (10-12 pt for body text)."
        })

    return findings


def _check_invisible_characters(text: str) -> List[Dict[str, Any]]:
    findings = []
    zw_matches = ZERO_WIDTH_RE.findall(text)
    if len(zw_matches) >= 3:
        findings.append({
            "id": "invisible_zero_width_chars",
            "severity": "high",
            "title": "Invisible / Zero-Width Characters Detected",
            "detail": f"Detected {len(zw_matches)} zero-width or formatting characters (e.g. U+200B, U+FEFF, U+200D).",
            "evidence": f"Found {len(zw_matches)} occurrences embedded within text.",
            "fix": "Remove zero-width spaces and hidden unicode control characters. ATS parsers sanitize or flag documents with embedded hidden unicode."
        })

    pad_matches = EXCESSIVE_PADDING_RE.findall(text)
    if pad_matches:
        findings.append({
            "id": "abnormal_whitespace_padding",
            "severity": "medium",
            "title": "Abnormal Whitespace Padding",
            "detail": "Detected long sequences of 25+ consecutive spaces or excessive empty lines.",
            "evidence": f"Longest gap: {len(max(pad_matches, key=len))} spaces/newlines.",
            "fix": "Use document margin and paragraph formatting rather than consecutive spaces to align text."
        })

    return findings


def _check_keyword_stuffing(text: str, skills_extracted: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    findings = []
    words = re.findall(r"\b[a-zA-Z]{4,}\b", text.lower())
    total_words = len(words)

    # 1. Frequency outliers
    STOPWORDS = {
        "with", "that", "this", "from", "have", "were", "been", "their", "which",
        "about", "into", "more", "other", "some", "such", "than", "then", "them",
        "these", "will", "would", "using", "used", "work", "year", "years", "page"
    }
    counts: Dict[str, int] = {}
    for w in words:
        if w not in STOPWORDS:
            counts[w] = counts.get(w, 0) + 1

    overused = []
    if total_words > 80:
        for w, c in counts.items():
            freq = c / total_words
            if c >= 10 and freq > 0.045:
                overused.append((w, c, round(freq * 100, 1)))

    if overused:
        overused.sort(key=lambda x: -x[1])
        top_w = overused[0]
        findings.append({
            "id": "keyword_stuffing_repetition",
            "severity": "medium",
            "title": "High Repetition / Keyword Density Outlier",
            "detail": f"The term '{top_w[0]}' appears {top_w[1]} times ({top_w[2]}% of all resume words), exceeding natural prose distribution.",
            "evidence": ", ".join(f"'{w}' ({c}x)" for w, c, _ in overused[:3]),
            "fix": f"Vary your vocabulary and describe outcomes rather than repeating '{top_w[0]}' excessively."
        })

    # 2. Skill dump (> 40 items)
    if len(skills_extracted) > 40:
        findings.append({
            "id": "skill_dump_excessive",
            "severity": "medium",
            "title": "Excessive Skill List (> 40 items)",
            "detail": f"Resume contains {len(skills_extracted)} listed skills, which triggers low-signal keyword dumping filters.",
            "evidence": f"Total extracted skills: {len(skills_extracted)}",
            "fix": "Curate your skills section to 15-25 high-conviction tools and frameworks tailored to the target role."
        })

    # 3. Unnaturally clustered technical buzzwords
    cluster_match = re.search(r"((?:[A-Za-z0-9#+.-]+,\s*){8,}[A-Za-z0-9#+.-]+)", text)
    if cluster_match and len(cluster_match.group(1).split(",")) >= 10:
        findings.append({
            "id": "clustered_keywords",
            "severity": "low",
            "title": "Dense Keyword Clustering",
            "detail": "Found a dense comma-separated list of 10+ technical buzzwords without supporting context.",
            "evidence": cluster_match.group(1)[:80] + "...",
            "fix": "Group skills logically into categories (e.g. Languages, Frameworks, Cloud) and embed key tools in project descriptions."
        })

    return findings


def _check_jd_copy_paste(resume_text: str, jd_text: Optional[str]) -> List[Dict[str, Any]]:
    findings = []
    if not jd_text or len(jd_text.strip()) < 100:
        return findings

    r_words = re.findall(r"\b[a-zA-Z0-9_-]+\b", resume_text.lower())
    j_words = re.findall(r"\b[a-zA-Z0-9_-]+\b", jd_text.lower())

    if len(r_words) < 20 or len(j_words) < 20:
        return findings

    # Build 6-grams
    def get_6grams(words):
        return set(tuple(words[i:i+6]) for i in range(len(words) - 5))

    r_6grams = get_6grams(r_words)
    j_6grams = get_6grams(j_words)

    if not r_6grams:
        return findings

    overlap = r_6grams & j_6grams
    share = len(overlap) / len(r_6grams)

    if share > 0.12: # Flag above 12%
        sample = " ".join(next(iter(overlap)))
        findings.append({
            "id": "verbatim_jd_copy_paste",
            "severity": "high",
            "title": "Verbatim Job Description Copy-Paste Detected",
            "detail": f"{round(share * 100, 1)}% of 6-word sequences in your resume appear verbatim in the target JD (threshold: 12%).",
            "evidence": f"Verbatim sequence example: '{sample}'",
            "fix": "Paraphrase your experience in your own voice. Copying verbatim clauses from the JD triggers automated plagiarism and integrity filters."
        })

    return findings


def _check_unsupported_claims(text: str, skills_extracted: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    findings = []
    if not skills_extracted:
        return findings

    # Look for Experience and Project sections
    text_lower = text.lower()
    exp_idx = re.search(r"\b(experience|work experience|employment|internships?)\b", text_lower)
    proj_idx = re.search(r"\b(projects|personal projects|technical projects)\b", text_lower)
    skills_idx = re.search(r"\b(skills|technical skills|competencies)\b", text_lower)

    # Extract body of experience/projects
    body_text = ""
    if exp_idx or proj_idx:
        start = min(m.start() for m in (exp_idx, proj_idx) if m)
        # cut before next major heading if skills is later
        end = skills_idx.start() if skills_idx and skills_idx.start() > start else len(text_lower)
        body_text = text_lower[start:end]
    else:
        body_text = text_lower

    unsupported = []
    for s in skills_extracted:
        name = s.get("skill_name", "").lower()
        if len(name) < 3:
            continue
        # Check if skill appears in body text outside of a bare list
        if not re.search(r"\b" + re.escape(name) + r"\b", body_text):
            unsupported.append(s.get("skill_name"))

    if len(skills_extracted) >= 8 and len(unsupported) >= 6 and (len(unsupported) / len(skills_extracted)) > 0.40:
        findings.append({
            "id": "unsupported_skill_claims",
            "severity": "low",
            "title": "Skills Lacking Project/Experience Context",
            "detail": f"{len(unsupported)} of your {len(skills_extracted)} listed skills have no accompanying mention in your Experience or Projects bullets.",
            "evidence": "Skills without evidence: " + ", ".join(unsupported[:5]) + ("..." if len(unsupported) > 5 else ""),
            "fix": "Anchor skills with evidence: add 1-2 bullet points demonstrating practical application and outcomes for each core skill."
        })

    return findings


def _check_parseability(text: str, pdf_bytes: Optional[bytes] = None, docx_bytes: Optional[bytes] = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    findings = []
    stats = {}

    char_count = len(text)
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)
    stats["char_count"] = char_count
    stats["word_count"] = word_count

    # 1. No text layer (scanned/raster)
    if char_count < 60:
        findings.append({
            "id": "scanned_no_text_layer",
            "severity": "high",
            "title": "Scanned Document / Missing Text Layer",
            "detail": f"Only {char_count} characters could be extracted. The document may be a scanned image or rasterized PDF.",
            "evidence": f"Extracted characters: {char_count}",
            "fix": "Export your resume directly from Word, Google Docs, or LaTeX as a text-based PDF rather than scanning a printed page."
        })
        stats["page_count"] = 1
        stats["detected_sections"] = []
        stats["has_contact_info"] = False
        return findings, stats

    # 2. Page count & table layout check
    page_count = 1
    has_tables = False
    if pdf_bytes:
        try:
            if pdfplumber:
                with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                    page_count = len(pdf.pages)
                    for page in pdf.pages:
                        tables = page.find_tables()
                        if len(tables) > 1:
                            has_tables = True
            else:
                reader = PdfReader(io.BytesIO(pdf_bytes))
                page_count = len(reader.pages)
        except Exception:
            pass
    elif docx_bytes:
        try:
            doc = Document(io.BytesIO(docx_bytes))
            if len(doc.tables) > 1:
                has_tables = True
        except Exception:
            pass

    stats["page_count"] = page_count

    if has_tables:
        findings.append({
            "id": "complex_table_layout",
            "severity": "low",
            "title": "Complex Table Layout Detected",
            "detail": "Multiple table structures were detected in the document layout.",
            "evidence": "Document contains embedded table bounding boxes.",
            "fix": "Use clean, single-column margins and standard text tabs instead of nested tables to avoid layout corruption in older ATS systems."
        })

    # 3. Missing standard sections
    text_lower = text.lower()
    detected_sections = []
    for sec_name, pattern in STANDARD_SECTIONS.items():
        if re.search(pattern, text_lower):
            detected_sections.append(sec_name.title())

    stats["detected_sections"] = detected_sections
    missing_core = []
    if "Education" not in detected_sections:
        missing_core.append("Education")
    if "Experience" not in detected_sections:
        missing_core.append("Experience")

    if missing_core:
        findings.append({
            "id": "missing_standard_sections",
            "severity": "medium",
            "title": f"Missing Core Section Header: {', '.join(missing_core)}",
            "detail": f"ATS parser could not identify standard section headers for: {', '.join(missing_core)}.",
            "evidence": f"Sections identified: {', '.join(detected_sections) if detected_sections else 'None'}",
            "fix": "Use standard, unambiguous section titles such as 'Work Experience' and 'Education'."
        })

    # 4. Contact info
    has_email = bool(EMAIL_RE.search(text))
    has_phone = bool(PHONE_RE.search(text))
    stats["has_contact_info"] = has_email and has_phone

    if not has_email or not has_phone:
        missing_contacts = []
        if not has_email:
            missing_contacts.append("email address")
        if not has_phone:
            missing_contacts.append("phone number")
        findings.append({
            "id": "missing_contact_details",
            "severity": "low",
            "title": f"Missing Contact Information ({', '.join(missing_contacts)})",
            "detail": f"Could not find a parseable {', '.join(missing_contacts)} in the document text.",
            "evidence": f"Email found: {has_email}, Phone found: {has_phone}",
            "fix": "Ensure your professional email and contact phone number are placed at the very top of page 1 as plain text."
        })

    return findings, stats


def ats_check(
    file_bytes: bytes,
    filename: str,
    text: str,
    jd_text: Optional[str] = None,
    skills_extracted: Optional[List[Dict[str, Any]]] = None
) -> Dict[str, Any]:
    """
    Performs comprehensive ATS integrity and risk screening on a resume.
    Returns: {risk_score: 0-100, level: 'Low'|'Medium'|'High', findings: [...], stats: {...}}
    """
    findings: List[Dict[str, Any]] = []
    fname_lower = filename.lower()
    skills = skills_extracted or []

    is_pdf = fname_lower.endswith(".pdf")
    is_docx = fname_lower.endswith((".docx", ".doc"))

    # 1. Hidden text
    if is_pdf:
        findings.extend(_check_hidden_pdf(file_bytes))
    elif is_docx:
        findings.extend(_check_hidden_docx(file_bytes))

    # 2. Invisible characters
    findings.extend(_check_invisible_characters(text))

    # 3. Keyword stuffing
    findings.extend(_check_keyword_stuffing(text, skills))

    # 4. JD copy-paste
    findings.extend(_check_jd_copy_paste(text, jd_text))

    # 5. Unsupported claims
    findings.extend(_check_unsupported_claims(text, skills))

    # 6. Parseability & Stats
    parse_findings, stats = _check_parseability(
        text,
        pdf_bytes=file_bytes if is_pdf else None,
        docx_bytes=file_bytes if is_docx else None
    )
    findings.extend(parse_findings)

    # Calculate weighted risk score
    # High = 25 pts, Medium = 12 pts, Low = 5 pts, Info = 0 pts (capped at 100)
    raw_score = sum(SEVERITY_WEIGHTS.get(f.get("severity", "low"), 0) for f in findings)
    risk_score = min(100, raw_score)

    # Risk level:
    # Low: score < 15 (clean or minor single low finding)
    # Medium: 15 <= score < 25 (moderate formatting/section issues)
    # High: score >= 25 (intentional cheat detection or multiple critical issues)
    has_high_finding = any(f.get("severity") == "high" for f in findings)
    if risk_score >= 25 or has_high_finding:
        level = "High"
    elif risk_score >= 15:
        level = "Medium"
    else:
        level = "Low"

    return {
        "risk_score": risk_score,
        "level": level,
        "findings": findings,
        "stats": stats
    }
