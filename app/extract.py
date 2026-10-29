import hashlib
import io
import json
import os
import re
from typing import List, Optional, Tuple

from docx import Document
from pydantic import BaseModel, Field, field_validator
from pypdf import PdfReader

from .config import CODES, LLM_PROVIDER, MODEL, SKILLSET_MAP
from .store import get_store

VALID_CODES = set(CODES)

# ---------------------------------------------------------------------------
# Pydantic Contract
# ---------------------------------------------------------------------------
class SkillItem(BaseModel):
    skill_name: str
    category_code: str
    evidence: str = ""
    confidence: str = "medium"

    @field_validator("category_code", mode="before")
    def clamp_category(cls, v):
        code = str(v).strip().upper() if v else "OTHER"
        return code if code in VALID_CODES else "OTHER"

    @field_validator("evidence", mode="before")
    def clean_evidence(cls, v):
        if not v:
            return ""
        clean = re.sub(r"\s+", " ", str(v)).strip()
        return clean[:120]

    @field_validator("confidence", mode="before")
    def clamp_confidence(cls, v):
        c = str(v).strip().lower() if v else "medium"
        return c if c in {"high", "medium", "low"} else "medium"

    @field_validator("skill_name", mode="before")
    def clean_skill_name(cls, v):
        return re.sub(r"\s+", " ", str(v)).strip()


class ExtractionContract(BaseModel):
    company: str = ""
    role: str = ""
    skills: List[SkillItem] = Field(default_factory=list)


def dedupe_skills(skills: List[SkillItem]) -> List[dict]:
    seen = set()
    deduped = []
    for item in skills:
        key = (item.skill_name.lower().strip(), item.category_code)
        if key not in seen:
            seen.add(key)
            deduped.append(item.model_dump())
    return deduped


# ---------------------------------------------------------------------------
# Text Extraction Helpers
# ---------------------------------------------------------------------------
def _extract_docx_text(src) -> str:
    """Extract complete text from DOCX including headers, paragraphs, and tables."""
    try:
        doc = Document(src)
    except Exception:
        return ""

    parts: List[str] = []

    # 1. Header paragraphs from all sections
    try:
        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    t = p.text.strip()
                    if t:
                        parts.append(t)
    except Exception:
        pass

    # 2. Main body paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)

    # 3. Tables (crucial for resume skills, experience & education grids)
    for table in doc.tables:
        for row in table.rows:
            row_cells: List[str] = []
            last_text = None
            for cell in row.cells:
                t = cell.text.strip()
                if t and t != last_text:
                    row_cells.append(t)
                    last_text = t
            if row_cells:
                parts.append(" | ".join(row_cells))

    # 4. Footer paragraphs from all sections
    try:
        for section in doc.sections:
            if section.footer:
                for p in section.footer.paragraphs:
                    t = p.text.strip()
                    if t:
                        parts.append(t)
    except Exception:
        pass

    return "\n".join(parts)


def _extract_doc_binary_text(raw_bytes: bytes) -> str:
    """Extract readable text streams from legacy Word 97-2003 (.doc) binary files."""
    # Try UTF-16LE text sequences (standard in WordDocument stream)
    utf16_matches = re.findall(rb'(?:[\x20-\x7E\t\n\r]\x00){3,}', raw_bytes)
    utf16_text = [m.decode('utf-16le', errors='ignore').strip() for m in utf16_matches if len(m) >= 6]

    # Try ASCII / Latin-1 text sequences
    ascii_matches = re.findall(rb'[\x20-\x7E\t\n\r]{4,}', raw_bytes)
    ascii_text = [m.decode('utf-8', errors='ignore').strip() for m in ascii_matches if len(m) >= 4]

    combined: List[str] = []
    seen = set()
    for chunk in utf16_text + ascii_text:
        chunk_clean = chunk.strip()
        if not chunk_clean:
            continue
        low = chunk_clean.lower()
        if any(bad in low for bad in ("times new roman", "calibri", "normal.dotm", "microsoft word", "root entry", "worddocument")):
            continue
        if len(chunk_clean) > 2 and chunk_clean not in seen:
            seen.add(chunk_clean)
            combined.append(chunk_clean)

    return "\n".join(combined)


def read_text(path_or_bytes, name: str = "") -> str:
    name = (name or str(path_or_bytes)).lower()
    raw_bytes: Optional[bytes] = None
    if isinstance(path_or_bytes, bytes):
        raw_bytes = path_or_bytes
        src = io.BytesIO(path_or_bytes)
    else:
        src = path_or_bytes
        if isinstance(src, str) and os.path.exists(src):
            try:
                with open(src, "rb") as f:
                    raw_bytes = f.read()
            except Exception:
                pass

    # DOCX format
    if name.endswith(".docx"):
        extracted = _extract_docx_text(src)
        if extracted:
            return extracted.replace("\x00", "")

    # Legacy DOC format
    if name.endswith(".doc"):
        # First attempt python-docx (many modern files are docx with .doc extension)
        extracted = _extract_docx_text(src)
        if extracted and len(extracted) > 50:
            return extracted.replace("\x00", "")
        # Fallback to binary OLE stream extraction
        if raw_bytes is None and isinstance(src, io.BytesIO):
            raw_bytes = src.getvalue()
        if raw_bytes:
            return _extract_doc_binary_text(raw_bytes).replace("\x00", "")
        return ""

    # Plain text / Markdown / JSON / CSV
    if name.endswith((".txt", ".text", ".md", ".json", ".csv")):
        if isinstance(src, io.BytesIO):
            return src.getvalue().decode("utf-8", errors="ignore").replace("\x00", "")
        if isinstance(path_or_bytes, bytes):
            return path_or_bytes.decode("utf-8", errors="ignore").replace("\x00", "")
        try:
            with open(src, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().replace("\x00", "")
        except Exception:
            return ""

    # PDF format
    try:
        reader = PdfReader(src)
        pdf_text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if pdf_text.strip():
            return pdf_text.replace("\x00", "")
    except Exception:
        pass

    # Secondary PDF fallback using pdfplumber if pypdf yielded sparse or no text
    if raw_bytes or isinstance(src, io.BytesIO):
        b = raw_bytes if raw_bytes is not None else src.getvalue()
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(b)) as pdf:
                plumb_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
                if plumb_text.strip():
                    return plumb_text.replace("\x00", "")
        except Exception:
            pass

    if isinstance(src, io.BytesIO):
        decoded = src.getvalue().decode("utf-8", errors="ignore")
        return decoded.replace("\x00", "")
    try:
        with open(src, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().replace("\x00", "")
    except Exception:
        return ""


# ---------------------------------------------------------------------------
# Keyword Fallback (Regex-based, false-positive resilient)
# ---------------------------------------------------------------------------
SKIP_RE = re.compile(
    r"\b(ctc|bonus|rsus?|stock (?:awards|grants)|insurance|full-time|part-time|shift allowance)\b",
    re.IGNORECASE,
)

KW = {
    "COD": (
        r"(?:\bc\+\+|\bjava\b|\bpython\b|\bjavascript\b|\btypescript\b|\bgolang\b|\bc#\b|\brust\b|\bcoding\b|\bprogramming\b)",
        "Programming & Languages",
    ),
    "DSA": (
        r"\b(data structures|algorithms?|competitive programming|problems solved|dsa|dynamic programming|trees?|graphs?)\b",
        "Data Structures & Algorithms",
    ),
    "OOD": (
        r"\b(object[- ]oriented|oop|design patterns?|solid principles?|factory pattern|strategy pattern)\b",
        "Object-Oriented Design",
    ),
    "SYSD": (
        r"\b(system design|distributed systems?|distributed architecture|horizontally[- ]scalable|scalability|scalable software|pub[- ]sub|microservices?|rate[- ]limiter|load balancer|high availability)\b",
        "System Design & Architecture",
    ),
    "SQL": (
        r"\b(sql|mysql|postgresql|postgres|database schemas?|query optimization|indexing|relational databases?|queries)\b",
        "Databases & SQL",
    ),
    "CLOUD": (
        r"\b(cloud|azure|aws|gcp|docker|kubernetes|containerization|serverless)\b",
        "Cloud & DevOps",
    ),
    "SWE": (
        r"\b(software engineering|code review|ci/cd|\bci\b|\bcd\b|technical debt|agile|scrum|git\b|sdlc|testing|production code|troubleshooting|helpdesk|support desk|ticketing)\b",
        "Software Engineering Practices",
    ),
    "NETW": (
        r"\b(computer networks?|networking fundamentals?|tcp/?ip|udp|dns|http/https?|rest apis?|latency|ip address|cisco)\b",
        "Computer Networking",
    ),
    "OS": (
        r"\b(operating systems?|\bos\b|linux|unix|multithreaded|multithreading|threads?|virtual memory|process scheduling|kernel)\b",
        "Operating Systems",
    ),
    "COMM": (
        r"\b(communication skills?|stakeholders?|collaboration|collaborative|cross[- ]functional|presentation skills?|present findings|volunteer coordinator|coordinating logistics)\b",
        "Communication & Collaboration",
    ),
    "OTHER": (
        r"\b(power bi|tableau|data visualization|dashboards?|bi tools?|reporting|telemetry)\b",
        "Data Visualization & Reporting",
    ),
}

APTI_TRIGGERS = [
    "aptitude",
    "quantitative",
    "probability",
    "data interpretation",
    "analytical",
    "numerical",
    "problem solving",
    "excel",
]

AI_TRIGGERS = [
    "machine learning",
    "deep learning",
    "neural network",
    "regression",
    "classification",
    "scikit",
    "scikit-learn",
    "overfitting",
    "bias-variance",
    "gradient descent",
    "predictive models",
    "statistical",
    "statistics",
    "a/b test",
    "a/b testing",
    "collaborative filtering",
    "recommendation engine",
    "precision/recall",
    "pandas",
    "ml",
]


def extract_title_info(text: str) -> Tuple[str, str]:
    """Extract company and role from JD title line."""
    lines = [re.sub(r"\s+", " ", l).strip() for l in text.split("\n") if l.strip()]
    if not lines:
        return "", ""
    first_few = " ".join(lines[:3])
    match = re.search(r"^(.*?)\s+[—–-]\s+(.*?)(?:\s+Role Category:|\s+ Bengaluru|\s+ Hyderabad|\s+ Chennai|\s+ Mumbai|\Z)", first_few)
    if match:
        comp = match.group(1).strip()
        role = match.group(2).strip()
        return comp, role
    if "—" in lines[0]:
        parts = lines[0].split("—", 1)
        return parts[0].strip(), parts[1].strip()
    return "", ""


def get_clean_blocks(text: str, kind: str = "jd") -> List[str]:
    lines = [l.strip() for l in text.split("\n")]
    blocks = []
    current = []
    headers = [
        "education", "technical skills", "projects", "hackathons", "certifications",
        "achievements", "preferred roles", "key responsibilities", "role overview",
        "what we", "nice to have", "benefits", "role category"
    ]
    for line in lines:
        if not line:
            if current:
                blocks.append(" ".join(current))
                current = []
        elif line.startswith("•") or line.startswith("-") or line.startswith("*"):
            if current:
                blocks.append(" ".join(current))
            current = [line.lstrip("•-* ").strip()]
        elif any(line.lower().startswith(h) for h in headers):
            if current:
                blocks.append(" ".join(current))
                current = []
        else:
            current.append(line)
    if current:
        blocks.append(" ".join(current))

    clean = []
    for i, b in enumerate(blocks):
        b_clean = re.sub(r"\s+", " ", b).strip()
        low = b_clean.lower()
        if not b_clean:
            continue
        if SKIP_RE.search(b_clean):
            continue
        if "radix talent match" in low or "page 1" in low or "page 2" in low:
            continue
        # Skip candidate header/taglines or JD title in evidence
        if i == 0 and (kind == "jd" or "fit for" in low or "generalist" in low or "all-rounder" in low):
            continue
        clean.append(b_clean)
    return clean


def select_best_evidence(hits: List[str]) -> str:
    for h in hits:
        if any(w in h.lower() for w in [
            "built", "designed", "contributed", "intern", "hackathon", "project",
            "shipped", "developed", "architect", "responsibilities", "course",
            "certif", "coordinator", "simulation", "experience", "lead", "manage"
        ]):
            return h
    return hits[0]


def fallback_extract(text: str, kind: str = "jd") -> dict:
    blocks = get_clean_blocks(text, kind)
    clean_full_text = " ".join(blocks)
    skills = []

    # Standard keyword extraction
    for code, (pat, skill_name) in KW.items():
        hits = [b for b in blocks if re.search(pat, b, re.I)]
        if hits:
            best_ev = select_best_evidence(hits)
            confidence = "high" if len(hits) >= 3 else "medium" if len(hits) == 2 else "low"
            # If resume and evidenced in applied projects/internships, elevate confidence
            if kind == "resume" and any(w in best_ev.lower() for w in ["built", "shipped", "intern", "contributed", "designed", "winner"]):
                confidence = "high"
            skills.append(
                {
                    "skill_name": skill_name,
                    "category_code": code,
                    "evidence": best_ev[:120],
                    "confidence": confidence,
                }
            )

    # APTI requires 2+ distinct triggers
    matched_apti = {t for t in APTI_TRIGGERS if re.search(r"\b" + re.escape(t) + r"\b", clean_full_text, re.I)}
    if len(matched_apti) >= 2 and not any(s["category_code"] == "APTI" for s in skills):
        apti_hits = [b for b in blocks if any(re.search(r"\b" + re.escape(t) + r"\b", b, re.I) for t in matched_apti)]
        if apti_hits:
            best_ev = select_best_evidence(apti_hits)
            skills.append(
                {
                    "skill_name": "Quantitative & Analytical Aptitude",
                    "category_code": "APTI",
                    "evidence": best_ev[:120],
                    "confidence": "high" if len(matched_apti) >= 3 else "medium",
                }
            )

    # AI requires 2+ distinct triggers
    matched_ai = {t for t in AI_TRIGGERS if re.search(r"\b" + re.escape(t) + r"\b", clean_full_text, re.I)}
    if len(matched_ai) >= 2 and not any(s["category_code"] == "AI" for s in skills):
        ai_hits = [b for b in blocks if any(re.search(r"\b" + re.escape(t) + r"\b", b, re.I) for t in matched_ai)]
        if ai_hits:
            best_ev = select_best_evidence(ai_hits)
            skills.append(
                {
                    "skill_name": "Machine Learning & Statistics",
                    "category_code": "AI",
                    "evidence": best_ev[:120],
                    "confidence": "high" if len(matched_ai) >= 3 else "medium",
                }
            )

    company, role = ("", "")
    if kind == "jd":
        company, role = extract_title_info(text)

    return {"company": company, "role": role, "skills": skills}


# ---------------------------------------------------------------------------
# LLM Extraction Prompt & Engine
# ---------------------------------------------------------------------------
PROMPT_TEMPLATE = """You are an expert technical talent assessment engine.
Extract skills from the provided {doc_type} text.
Return ONLY a valid JSON object matching the contract below. Do NOT output any markdown ticks (```) or preamble.

JSON Contract:
{{
  "company": "{company}",
  "role": "{role}",
  "skills": [
    {{
      "skill_name": "Standardized skill name",
      "category_code": "COD|DSA|APTI|COMM|OOD|AI|SQL|SYSD|CLOUD|SWE|NETW|OS|OTHER",
      "evidence": "Short exact quote under 120 chars from text (NEVER include salary, CTC, bonus, RSUs, insurance, benefits)",
      "confidence": "high|medium|low"
    }}
  ]
}}

Category Code Definitions:
- COD: programming/coding ability and languages
- DSA: data structures and algorithms
- APTI: quantitative and logical aptitude
- COMM: communication
- OOD: object-oriented design and patterns
- AI: machine learning, statistics, AI
- SQL: databases and SQL
- SYSD: system design and architecture
- CLOUD: cloud and DevOps
- SWE: testing, code review, CI/CD, agile, engineering practice
- NETW: networking
- OS: operating systems
- OTHER: other specialized technical or domain skills

Rules:
{rules}

TEXT:
\"\"\"{text}\"\"\"
"""

JD_RULES = """- Extract company and role from the JD title line (e.g. for 'Google LLC — Data Scientist', company is 'Google LLC' and role is 'Data Scientist').
- Extract ONLY skills the job actually requires (focus on 'Key Responsibilities' and 'What We're Looking For').
- One entry per distinct skill, max ~15 skills total.
- evidence: short quote under 120 chars. Never take evidence containing CTC, bonus, RSUs, insurance, or Full-Time.
- confidence: how strongly the JD stresses it (high = explicitly required/emphasised, medium = mentioned, low = nice-to-have)."""

RESUME_RULES = """- For resumes, judge confidence from projects, internships and achievements, not just the skills list:
  * high = skill is listed AND demonstrated in a project or internship
  * medium = skill is only listed in technical skills or coursework
  * low = vague or minimal mention
- Include education/certifications evidence where relevant.
- evidence: short quote under 120 chars from projects/internships/education.
- Leave company and role as empty strings."""


def call_llm(prompt: str, provider: str, model: str) -> str:
    if provider == "gemini":
        from google import genai
        from google.genai import types

        api_key = os.getenv("GEMINI_API_KEY")
        if not api_key:
            raise ValueError("GEMINI_API_KEY is not set")
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model=model,
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                temperature=0.0,
            ),
        )
        return response.text or ""

    elif provider == "groq":
        from openai import OpenAI

        api_key = os.getenv("GROQ_API_KEY") or os.getenv("OPENAI_API_KEY")
        if not api_key:
            raise ValueError("GROQ_API_KEY is not set")
        base_url = os.getenv("GROQ_BASE_URL", "https://api.groq.com/openai/v1")
        client = OpenAI(base_url=base_url, api_key=api_key)
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.0,
            response_format={"type": "json_object"},
        )
        return response.choices[0].message.content or ""

    raise ValueError(f"Unknown LLM provider: {provider}")


def parse_and_validate(raw_json: str) -> dict:
    s = raw_json.strip()
    if s.startswith("```"):
        s = re.sub(r"^```(?:json)?\s*", "", s)
        s = re.sub(r"\s*```$", "", s)
    start = s.find("{")
    end = s.rfind("}")
    if start != -1 and end != -1:
        s = s[start : end + 1]

    data = json.loads(s)
    contract = ExtractionContract.model_validate(data)
    deduped = dedupe_skills(contract.skills)
    return {
        "company": contract.company.strip(),
        "role": contract.role.strip(),
        "skills": deduped,
    }


def llm_extract(text: str, kind: str) -> dict:
    provider = os.getenv("LLM_PROVIDER", LLM_PROVIDER).lower()
    model = os.getenv("RADIX_MODEL", MODEL)
    doc_type = "Job Description" if kind == "jd" else "Resume"
    rules = JD_RULES if kind == "jd" else RESUME_RULES
    company_hint, role_hint = extract_title_info(text) if kind == "jd" else ("", "")
    prompt = PROMPT_TEMPLATE.format(
        doc_type=doc_type,
        company=company_hint,
        role=role_hint,
        rules=rules,
        text=text[:12000],
    )

    last_err = None
    for attempt in range(2):
        try:
            raw = call_llm(prompt, provider, model)
            return parse_and_validate(raw)
        except Exception as e:
            last_err = e
            if attempt < 1:
                import time
                time.sleep(1)
                continue
    raise last_err or RuntimeError("LLM extraction failed after retry")


# ---------------------------------------------------------------------------
# Cache & Primary Extraction
# ---------------------------------------------------------------------------
def _get_cache_path(raw_bytes: bytes, kind: str, model: str) -> str:
    h = hashlib.sha256()
    h.update(raw_bytes)
    h.update(kind.encode("utf-8"))
    h.update(model.encode("utf-8"))
    cache_dir = os.path.join("data", "cache")
    os.makedirs(cache_dir, exist_ok=True)
    return os.path.join(cache_dir, f"{h.hexdigest()}.json")


def is_cached(raw_bytes: bytes, kind: str) -> bool:
    model = os.getenv("RADIX_MODEL", MODEL)
    h = hashlib.sha256(raw_bytes + kind.encode("utf-8") + model.encode("utf-8")).hexdigest()
    bundled_path = os.path.join("data", "cache", f"{h}.json")
    if os.path.exists(bundled_path):
        return True
    store, _ = get_store()
    return store.get("extract", h) is not None


def is_llm_configured() -> bool:
    provider = os.getenv("LLM_PROVIDER", LLM_PROVIDER).lower()
    if provider == "gemini":
        return bool(os.getenv("GEMINI_API_KEY"))
    if provider == "groq":
        return bool(os.getenv("GROQ_API_KEY") or os.getenv("OPENAI_API_KEY"))
    return False


def extract(path_or_bytes, kind: str, name: str = "", use_cache: bool = True) -> dict:
    """kind: 'jd' | 'resume'. Returns shared-contract dict."""
    # Obtain raw bytes for hashing and text extraction
    if isinstance(path_or_bytes, bytes):
        raw_bytes = path_or_bytes
    else:
        try:
            with open(path_or_bytes, "rb") as f:
                raw_bytes = f.read()
        except Exception:
            raw_bytes = str(path_or_bytes).encode("utf-8")

    text = read_text(path_or_bytes, name)
    provider = os.getenv("LLM_PROVIDER", LLM_PROVIDER).lower()
    model = os.getenv("RADIX_MODEL", MODEL)
    
    h = hashlib.sha256()
    h.update(raw_bytes)
    h.update(kind.encode("utf-8"))
    h.update(model.encode("utf-8"))
    cache_key = h.hexdigest()

    # 1. Check cache first if configured and caching is enabled
    has_key = is_llm_configured()
    if use_cache and not os.getenv("RADIX_NO_CACHE"):
        # Bundled read-only seed/cache/ or data/cache/ first
        bundled_path = os.path.join("seed", "cache", f"{cache_key}.json")
        if not os.path.exists(bundled_path):
            bundled_path = os.path.join("data", "cache", f"{cache_key}.json")
        if os.path.exists(bundled_path):

            try:
                with open(bundled_path, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                if cached.get("skills"):
                    return {
                        "source_type": kind,
                        "source_file": name or str(path_or_bytes),
                        "company": cached.get("company", ""),
                        "role": cached.get("role", ""),
                        "skills": cached["skills"],
                        "_text": text if kind == "resume" else None,
                    }
            except Exception:
                pass

        # Check store (Supabase or Local)
        store, _ = get_store()
        cached = store.get("extract", cache_key)
        if cached and isinstance(cached, dict) and cached.get("skills"):
            return {
                "source_type": kind,
                "source_file": name or str(path_or_bytes),
                "company": cached.get("company", ""),
                "role": cached.get("role", ""),
                "skills": cached["skills"],
                "_text": text if kind == "resume" else None,
            }

    # 2. Try LLM if configured
    res = None
    if has_key:
        try:
            res = llm_extract(text, kind)
            # Save valid extraction to store
            if use_cache and res and res.get("skills"):
                store, _ = get_store()
                store.set("extract", cache_key, res)
        except Exception as e:
            print(f"[EXTRACT ERROR] LLM failed for {name or str(path_or_bytes)}: {type(e).__name__}: {e}")
            if os.getenv("RADIX_STRICT_LLM"):
                raise e

    # 3. Fallback extraction if no LLM key or LLM extraction failed
    if not res or not res.get("skills"):
        res = fallback_extract(text, kind)

    return {
        "source_type": kind,
        "source_file": name or str(path_or_bytes),
        "company": res.get("company", ""),
        "role": res.get("role", ""),
        "skills": res["skills"],
        "_text": text if kind == "resume" else None,
    }
